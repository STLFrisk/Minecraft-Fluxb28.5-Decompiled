#!/usr/bin/env python3
"""Render the delivery report markdown into the delivered .docx.

Handles exactly the constructs the report uses: level-1..4 headings, pipe tables,
fenced code blocks, blockquotes, bullet/numbered lists, and bold/italic/code spans.
Chinese text needs an explicit w:eastAsia font, which run.font.name alone does not
establish, so every run gets the eastAsia attribute set.
"""
import io
import re
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

SRC = sys.argv[1]
DST = sys.argv[2]

HEAD = re.compile(r'^(#{1,4})\s+(.*)$')
BULLET = re.compile(r'^\s*[-*]\s+(.*)$')
NUMBER = re.compile(r'^\s*(\d+)\.\s+(.*)$')
TABLE_SEP = re.compile(r'^\s*\|[\s:|-]+\|\s*$')


def setfont(run, name='Consolas'):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = rpr.makeelement(qn('w:rFonts'), {})
        rpr.append(rf)
    rf.set(qn('w:eastAsia'), '微软雅黑')
    rf.set(qn('w:ascii'), name)
    rf.set(qn('w:hAnsi'), name)


def add_rich(par, text):
    """Split on **bold**, `code`, *italic* and emit runs with formatting."""
    for part in re.split(r'(\*\*[^*]+\*\*|`[^`]+`|\*[^*]+\*)', text):
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            r = par.add_run(part[2:-2])
            r.bold = True
        elif part.startswith('`') and part.endswith('`'):
            r = par.add_run(part[1:-1])
            setfont(r, 'Consolas')
            r.font.size = Pt(9)
        elif part.startswith('*') and part.endswith('*') and len(part) > 2:
            r = par.add_run(part[1:-1])
            r.italic = True
        else:
            r = par.add_run(part)
        if r.font.name is None:
            setfont(r, '微软雅黑')


def cells(line):
    return [c.strip() for c in line.strip().strip('|').split('|')]


def main():
    lines = io.open(SRC, encoding='utf-8').read().split('\n')
    doc = Document()
    st = doc.styles['Normal']
    st.font.name = '微软雅黑'
    st.font.size = Pt(10)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')

    i = 0
    while i < len(lines):
        line = lines[i]

        fence = line.strip().startswith('```')
        if fence:
            i += 1
            buf = []
            while i < len(lines) and not lines[i].strip().startswith('```'):
                buf.append(lines[i])
                i += 1
            i += 1
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.left_indent = Pt(10)
            r = p.add_run('\n'.join(buf))
            setfont(r, 'Consolas')
            r.font.size = Pt(8.5)
            continue

        if line.startswith('|') and i + 1 < len(lines) and TABLE_SEP.match(lines[i + 1]):
            head = cells(line)
            i += 2
            rows = []
            while i < len(lines) and lines[i].startswith('|'):
                rows.append(cells(lines[i]))
                i += 1
            t = doc.add_table(rows=1, cols=len(head))
            t.style = 'Light Grid Accent 1'
            for k, h in enumerate(head):
                cell = t.rows[0].cells[k]
                cell.text = ''
                add_rich(cell.paragraphs[0], h)
                for r in cell.paragraphs[0].runs:
                    r.bold = True
                    r.font.size = Pt(9)
            for row in rows:
                cs = t.add_row().cells
                for k in range(len(head)):
                    cs[k].text = ''
                    add_rich(cs[k].paragraphs[0], row[k] if k < len(row) else '')
                    for r in cs[k].paragraphs[0].runs:
                        r.font.size = Pt(9)
            doc.add_paragraph()
            continue

        m = HEAD.match(line)
        if m:
            lvl = len(m.group(1))
            h = doc.add_heading(level=min(lvl, 4))
            add_rich(h, m.group(2))
            for r in h.runs:
                setfont(r, '微软雅黑')
            i += 1
            continue

        if line.strip().startswith('>'):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Pt(14)
            add_rich(p, line.strip().lstrip('>').strip())
            for r in p.runs:
                r.italic = True
                r.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
            i += 1
            continue

        if line.strip() == '---':
            doc.add_paragraph()
            i += 1
            continue

        b = BULLET.match(line)
        if b:
            p = doc.add_paragraph(style='List Bullet')
            add_rich(p, b.group(1))
            i += 1
            continue

        n = NUMBER.match(line)
        if n:
            p = doc.add_paragraph(style='List Number')
            add_rich(p, n.group(2))
            i += 1
            continue

        if not line.strip():
            i += 1
            continue

        p = doc.add_paragraph()
        add_rich(p, line)
        i += 1

    doc.save(DST)
    print('  %s  %d 段落  %d 表格' % (DST, len(doc.paragraphs), len(doc.tables)))


if __name__ == '__main__':
    main()
